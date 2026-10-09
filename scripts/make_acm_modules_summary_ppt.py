#!/usr/bin/env python3
"""Build a ten-slide, monochrome ACM summary with embedded scientific videos.

The archived 4,000-step run is the sole source of displayed results. The active
10,000-step run directory is not read or modified.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import re
import subprocess
from pathlib import Path

import h5py
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/presentations"
ASSETS = OUT / "assets/acm_modules_summary"
MEDIA = ASSETS / "media"
ARCHIVE = Path(
    "/mnt/ssd-SATARAID5/home/mrz/.local/share/Trash/files/"
    "acm2D_CylinderRe3900_BDF2_LUSGS"
)
CASE = ROOT / "cases/acm/2D/acm2D_cylinder_Re3900_laminar_BDF2_LUSGS.json"
PPTX = OUT / "ACM_Modules_Derivation_Results_Simple_CN.pptx"
FONT = "Noto Sans CJK SC"
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)


def read_config() -> dict:
    text = CASE.read_text(encoding="utf-8")
    return json.loads(re.sub(r"(?m)^\s*//.*$", "", text))


def audit_archived_results() -> dict:
    paths = sorted(ARCHIVE.glob("CylinderA1_Re3900_laminar_BDF2_LUSGS_*.vtkhdf"))
    steps = [int(path.stem.rsplit("_", 1)[1]) for path in paths]
    if steps != list(range(0, 4001, 100)):
        raise ValueError("The presentation requires the complete archived 4,000-step run")
    extrema = {key: [float("inf"), -float("inf")] for key in ("u", "v", "speed", "p")}
    final = {}
    for path in paths:
        with h5py.File(path, "r") as handle:
            velocity = handle["VTKHDF/CellData/Velocity"][:]
            pressure = handle["VTKHDF/CellData/Pressure"][:]
        if velocity.shape != (23791, 3) or pressure.shape != (23791,):
            raise ValueError(f"Unexpected field dimensions: {path}")
        fields = {"u": velocity[:, 0], "v": velocity[:, 1],
                  "speed": np.linalg.norm(velocity[:, :2], axis=1), "p": pressure}
        for key, field in fields.items():
            if not np.isfinite(field).all():
                raise ValueError(f"Non-finite {key} in {path}")
            bounds = [float(field.min()), float(field.max())]
            extrema[key] = [min(extrema[key][0], bounds[0]), max(extrema[key][1], bounds[1])]
            if path == paths[-1]:
                final[key] = bounds
    log = (ARCHIVE / "run_16mpi.log").read_text(encoding="utf-8")
    pattern = re.compile(
        r"ACM physical step\s+(\d+)\s+time=([\d.eE+-]+)\s+BDF(\d)"
        r"\s+residual\s+([\d.eE+-]+)\s+->\s+([\d.eE+-]+)"
        r".*?inner=(\d+)\s+converged=(\d+)"
    )
    rows = [[float(value) for value in match.groups()] for match in pattern.finditer(log)]
    if len(rows) != 4000 or int(rows[-1][0]) != 4000:
        raise ValueError("Archived physical-step log is incomplete")
    if "ACM flow field written at step 4000" not in log:
        raise ValueError("Missing final output completion record")
    history = np.asarray(rows)
    media = {}
    for field in ("density", "velocity", "u", "pressure"):
        video = MEDIA / f"CylinderA1_Re3900_BDF2_{field}_contours_steps_0000_4000.mp4"
        if not video.is_file() or not video.with_suffix(".png").is_file():
            raise FileNotFoundError(video)
        media[field] = {"name": video.name, "bytes": video.stat().st_size,
                        "sha256": hashlib.sha256(video.read_bytes()).hexdigest()}
    audit = {
        "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
        "data_source": str(ARCHIVE), "config_source": str(CASE),
        "physical_steps": 4000, "snapshot_count": len(paths), "cell_count": 23791,
        "final_time": 40.0, "snapshot_interval": 100, "mpi_ranks": 16,
        "rho0": 1.0, "density_is_config_constant": True,
        "final_cell_extrema": final, "global_snapshot_extrema": extrema,
        "last_initial_defect": history[-1, 3], "last_final_defect": history[-1, 4],
        "converged_step_count": int(history[:, 6].sum()),
        "all_iterations_equal_15": bool(np.all(history[:, 5] == 15)), "media": media,
    }
    (ASSETS / "result_audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    fig, ax = plt.subplots(figsize=(6.3, 3.85), dpi=190)
    ax.semilogy(history[:, 0], history[:, 3], color="black", linestyle="--",
                linewidth=0.8, label="Before inner iterations")
    ax.semilogy(history[:, 0], history[:, 4], color="black", linewidth=1.0,
                label="After inner iterations")
    ax.set(xlabel="Physical step", ylabel="Global RMS defect", xlim=(0, 4000))
    ax.legend(frameon=False, fontsize=9, loc="upper right")
    ax.tick_params(labelsize=9)
    fig.tight_layout()
    fig.savefig(ASSETS / "residual_history.png", facecolor="white")
    plt.close(fig)
    return audit


def render_equations() -> None:
    subprocess.run(
        ["xelatex", "-interaction=nonstopmode", "-halt-on-error", "equations.tex"],
        cwd=ASSETS, check=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
    )
    subprocess.run(
        ["pdftoppm", "-png", "-r", "200", str(ASSETS / "equations.pdf"),
         str(ASSETS / "eq")], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
    )
    for path in sorted(ASSETS.glob("eq-*.png")):
        with Image.open(path) as source:
            rgb = source.convert("RGB")
            bounds = ImageChops.difference(rgb, Image.new("RGB", rgb.size, "white")).getbbox()
            if bounds is None:
                raise ValueError(f"Blank equation page: {path}")
            crop = rgb.crop(bounds)
            padded = Image.new("RGB", (crop.width + 16, crop.height + 16), "white")
            padded.paste(crop, (8, 8))
            padded.save(path)


def add_text(slide, text, x, y, w, h, size=20, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.015)
    frame.margin_top = frame.margin_bottom = Inches(0.01)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(9)
        paragraph.line_spacing = 1.13
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = BLACK
        props = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            face = OxmlElement(tag)
            face.set("typeface", FONT)
            props.append(face)
    return shape


def add_slide(prs, title, notes, source):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_text(slide, title, 0.55, 0.35, 12.2, 0.66, size=28, bold=True)
    add_text(slide, source, 0.57, 7.10, 11.5, 0.22, size=10)
    add_text(slide, f"{len(prs.slides)}/10", 12.12, 7.10, 0.65, 0.22,
             size=10, align=PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_table(slide, rows, x, y, widths, row_heights, size=18):
    table = slide.shapes.add_table(
        len(rows), len(rows[0]), Inches(x), Inches(y),
        Inches(sum(widths)), Inches(sum(row_heights)),
    ).table
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    for index, height in enumerate(row_heights):
        table.rows[index].height = Inches(height)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.10)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            for side in ("lnL", "lnR", "lnT", "lnB"):
                edge = OxmlElement(f"a:{side}")
                edge.set("w", "5000")
                fill = OxmlElement("a:solidFill")
                color = OxmlElement("a:srgbClr")
                color.set("val", "000000")
                fill.append(color)
                edge.append(fill)
                cell._tc.get_or_add_tcPr().append(edge)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = FONT
                    run.font.size = Pt(size)
                    run.font.bold = i == 0
                    run.font.color.rgb = BLACK
                    props = run._r.get_or_add_rPr()
                    face = OxmlElement("a:ea")
                    face.set("typeface", FONT)
                    props.append(face)
    return table


def fit_image(slide, path, x, y, w, h):
    with Image.open(path) as source:
        ratio = source.width / source.height
    if w / h > ratio:
        width, height = h * ratio, h
    else:
        width, height = w, w / ratio
    return slide.shapes.add_picture(
        str(path), Inches(x + (w - width) / 2), Inches(y + (h - height) / 2),
        width=Inches(width), height=Inches(height),
    )


def equation(slide, number, x, y, w, h):
    return fit_image(slide, ASSETS / f"eq-{number:02d}.png", x, y, w, h)


def movie(slide, field, x, y, w=6.00):
    name = f"CylinderA1_Re3900_BDF2_{field}_contours_steps_0000_4000"
    return slide.shapes.add_movie(
        str(MEDIA / f"{name}.mp4"), Inches(x), Inches(y), Inches(w), Inches(w * 9 / 16),
        poster_frame_image=str(MEDIA / f"{name}.png"), mime_type="video/mp4",
    )


def sources(*items):
    return "\n".join(str(ROOT / item) for item in items)


def build_deck(config, audit):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = "常密度 ACM：模块、推导与圆柱绕流结果"
    prs.core_properties.subject = "DNDSR ACM implementation, numerical methods and verified case data"
    prs.core_properties.author = "DNDSR"
    prs.core_properties.keywords = "ACM, BDF2, CFV, CWBAP, Re3900, cylinder"

    # 1. Completed modules, with explicit implementation/validation boundaries.
    slide = add_slide(
        prs, "常密度 ACM：已完成的模块", sources(
            "src/ACM/README.md", "src/ACM/ACM.hpp:20", "src/ACM/ACMTurbulence.hpp:27",
            "test/cpp/ACM", "build/Testing/Temporary/LastTest.log",
        ) + "\n本汇报依据当前源码。测试源码存在不代表本次重新运行了全部测试。",
        "依据：当前 ACM 源码与测试记录；不采用旧审计中已过时的缺口结论",
    )
    add_table(slide, [
        ["模块", "已实现内容"],
        ["方程与通量", "二维 / 三维 [u,v,w,p]；Turkel Γ；Roe / Rusanov；牛顿黏性应力"],
        ["空间离散", "有限体积面求积；FirstOrder、GreenGauss、Variational"],
        ["限制器", "LocalExtrema、WBAP、CWBAP；ACM 特征空间变换"],
        ["时间与线性求解", "SSPRK3；隐式 Euler；BDF2 双时间；块 Jacobi / LU-SGS / GMRES"],
        ["工程化", "CGNS、METIS、MPI 幽灵层、周期平移、JSON、VTK-HDF 输出"],
        ["湍流扩展", "SA、Wilcox k-ω、SST k-ω、Realizable k-ε；分离式输运"],
    ], 0.58, 1.25, [2.12, 10.06], [0.46] + [0.59] * 6, size=18)
    add_text(slide, "验证：测试源码含 35 个核心 / 时间 / 湍流 / 并行用例；历史 MPI 1/2/4/8 进程测试通过。",
             0.61, 5.56, 12.04, 0.64, size=17)
    add_text(slide, "边界：尚无变密度、重力分层、自适应 β²；BDF2 仅接通 Laminar，重启历史尚未接通。",
             0.61, 6.27, 12.04, 0.54, size=17)

    # 2. Incompressible physics -> pseudo-time pressure constraint.
    slide = add_slide(
        prs, "推导 1：不可压方程 → 人工伪时间", sources(
            "src/ACM/ACM.cpp:133", "src/ACM/ACM.cpp:163", "src/ACM/ACM.cpp:645",
            "src/ACM/ACMSettings.hpp:27",
        ), "依据：ACM.cpp；状态中的 p 为物理压力变量，不是密度或能量",
    )
    add_text(slide, "假设：ρ=ρ₀ 为常数；当前算例无重力、无能量方程；二维时 w=0。",
             0.65, 1.17, 12.0, 0.50, size=20)
    equation(slide, 1, 0.74, 1.92, 11.82, 0.65)
    equation(slide, 2, 1.02, 2.93, 11.23, 0.62)
    add_text(slide, "压力没有独立物理演化方程；引入人工伪时间 τ，松弛不可压约束：",
             0.65, 4.02, 12.0, 0.56, size=20)
    equation(slide, 3, 0.96, 4.84, 11.45, 0.70)
    add_text(slide, "β² 控制数值上的人工波速。伪时间收敛后恢复 ∇·u=0，不引入真实可压缩密度波。",
             0.65, 6.09, 12.0, 0.73, size=20)

    # 3. Gamma matrix and actual coefficient settings.
    slide = add_slide(
        prs, "推导 2：Turkel 耦合与双时间方程", sources(
            "src/ACM/ACM.cpp:163", "src/ACM/ACM.cpp:175", "src/ACM/ACMBDF2.cpp:55",
            "cases/acm/2D/acm2D_cylinder_Re3900_laminar_BDF2_LUSGS.json:17",
        ), "依据：GammaLocal / GammaInvLocal / BDF2PhysicalMassMatrix",
    )
    add_text(slide, "将压力伪时间导数耦合到动量方程：γₜ=1+α。",
             0.65, 1.18, 12.0, 0.52, size=21)
    equation(slide, 4, 0.70, 1.95, 7.08, 2.02)
    add_table(slide, [["当前系数", "数值"], ["β² / α", "4.0 / 0.5"],
                      ["γₜ / ρ₀", "1.5 / 1.0"], ["1/β²", "0.25"],
                      ["γₜ/β²", "0.375"]],
              8.21, 1.91, [2.26, 2.10], [0.43] * 5, size=19)
    equation(slide, 5, 0.98, 4.39, 11.35, 0.72)
    add_text(slide, "R(Q) 是有限体积空间残差；t 是物理时间，τ 是每个物理步内的人工迭代时间。\n"
             "M 的压力分量为 0：BDF2 只对速度做物理时间差分。人工波速标度 √(β²/ρ₀)=2。",
             0.65, 5.66, 12.0, 1.11, size=20)

    # 4. Flux, characteristic speeds and finite-volume residual.
    slide = add_slide(
        prs, "推导 3：特征波 → 有限体积残差", sources(
            "src/ACM/ACM.cpp:133", "src/ACM/ACM.cpp:208", "src/ACM/ACM.cpp:505",
            "src/ACM/ACMEvaluator.hxx:443",
        ), "依据：法向物理通量、Γ⁻¹Aₙ 特征系统及 ACMEvaluator 残差组装",
    )
    equation(slide, 6, 0.98, 1.28, 11.35, 1.38)
    equation(slide, 7, 0.70, 3.07, 11.93, 0.64)
    equation(slide, 8, 1.90, 4.12, 9.53, 0.77)
    add_text(slide, "对流通量：Roe / Rusanov；Roe 使用熵修正，当前 entropyFixRatio=0.05。\n"
             "黏性通量：修正面梯度 + 面求积；内部面通量对相邻单元等量反号。\n"
             "流程：单元均值 → 重构 → 限制 → 面状态 / 梯度 → 数值通量 → R(Q)。",
             0.65, 5.42, 12.0, 1.31, size=19)

    # 5. Reconstruction and limiter implementation, exact degree language.
    slide = add_slide(
        prs, "空间模块：重构方法与限制器", sources(
            "src/ACM/ACMConfig.hpp:30", "src/ACM/ACMEvaluator.hxx:156",
            "src/ACM/ACMEvaluator.hxx:263", "src/ACM/ACMEvaluator.hxx:357",
            "src/CFV/VRSettings.hpp:175",
        ), "依据：ACM 直接枚举与 CFV 变分重构；精度阶数仍需网格收敛验证",
    )
    equation(slide, 9, 1.0, 1.19, 11.33, 0.70)
    add_text(slide, "重构路径", 0.65, 2.09, 5.8, 0.42, size=21, bold=True)
    add_text(slide, "限制器路径", 6.91, 2.09, 5.8, 0.42, size=21, bold=True)
    add_table(slide, [["方法", "实现"], ["FirstOrder", "单元分片常数"],
                      ["GreenGauss", "直接梯度 / 线性重构"],
                      ["Variational", "复用 CFV；可设多项式次数"]],
              0.64, 2.65, [2.08, 3.68], [0.43, 0.60, 0.60, 0.79], size=17.5)
    add_table(slide, [["方法", "实现"], ["LocalExtrema", "邻域极值约束，θ∈[0,1]"],
                      ["WBAP", "加权多项式限制"],
                      ["CWBAP", "逐阶加权多项式限制"]],
              6.91, 2.65, [2.00, 3.78], [0.43, 0.60, 0.60, 0.79], size=17.5)
    add_text(slide, "当前算例：Variational + maxOrder=2（二次多项式）+ CWBAP；每次残差做 3 次重构迭代。\n"
             "WBAP / CWBAP 只配合 Variational，并使用 ACM 四变量特征变换；二维 / 三维均有路径。\n"
             "二次多项式在光滑充分收敛条件下具备名义三阶潜力，不等于本算例已验证三阶精度。",
             0.65, 5.38, 12.0, 1.40, size=18)

    # 6. All time modules plus compact BDF2 derivation.
    slide = add_slide(
        prs, "时间模块：稳态伪时间与 BDF2 物理时间", sources(
            "src/ACM/ACMTime.hpp:27", "src/ACM/ACMBDF2.cpp:47",
            "src/ACM/ACMSolver.hxx:355", "src/ACM/ACMSolver.hxx:509",
        ) + "\n完整配置名称：ExplicitSSPRK3、ImplicitEulerBlockJacobi、ImplicitEulerLUSGS、"
        "ImplicitEulerGMRES、BDF2DualTimeLUSGS、BDF2DualTimeGMRES。",
        "依据：ACMTime、ACMBDF2、分布式 LU-SGS / GMRES 驱动",
    )
    add_table(slide, [["模式", "已实现格式 / 求解器"],
                      ["稳态伪时间", "SSPRK3；隐式 Euler + 块 Jacobi / LU-SGS / GMRES"],
                      ["非定常物理时间", "BDF2 双时间 + LU-SGS 或 GMRES；首步自动后向 Euler"]],
              0.64, 1.23, [2.30, 9.75], [0.42, 0.56, 0.56], size=18)
    equation(slide, 10, 1.60, 3.03, 10.15, 0.68)
    equation(slide, 12, 0.74, 4.07, 11.86, 0.71)
    add_text(slide, "Dᵏ = 空间残差 − BDF 物理时间项；首步 a₀=1，之后 a₀=3/2；M=diag(1,1,1,0)。\n"
             "隐式算子使用冻结的一阶近似面雅可比；非线性残差仍保留高阶重构，并非完整高阶 Newton。\n"
             "当前 BDF2：固定 Δt、仅 Laminar。达到内迭代上限仍提交该物理步，收敛标志另行报告。",
             0.65, 5.31, 12.0, 1.49, size=18)

    # 7. Shared physics and explicitly separated historical/current schedules.
    t = config["timeMarchSettings"]
    o = config["outputSettings"]
    slide = add_slide(
        prs, "二维圆柱 Re=3900：算例参数与运行版本", sources(
            "cases/acm/2D/acm2D_cylinder_Re3900_laminar_BDF2_LUSGS.json",
            "data/mesh/CylinderA1.cgns",
        ) + "\n已完成结果来源：" + str(ARCHIVE) + "\n当前配置可能由用户运行中；本PPT没有启动或重跑求解器。",
        "当前 JSON 与归档 4000 步结果分列；后两页视频只使用归档结果",
    )
    add_text(slide, "无量纲物理参数与边界", 0.64, 1.18, 5.8, 0.44, size=21, bold=True)
    add_text(slide, "完成结果 / 当前重算配置", 6.90, 1.18, 5.8, 0.44, size=21, bold=True)
    add_table(slide, [["参数", "设置"], ["D / U∞ / ρ₀", "1 / 1 / 1"],
                      ["Re = ρ₀U∞D / μ", "3900"], ["动力黏度 μ", "2.56410256×10⁻⁴"],
                      ["远场压力 / 重力", "p∞=0（表压）/ 无"], ["β² / α", "4 / 0.5（固定）"],
                      ["边界", "WALL 无滑移；FAR 特征远场"]],
              0.64, 1.84, [2.33, 3.46], [0.43] + [0.48] * 6, size=17)
    add_table(slide, [["参数", "已完成", "当前 JSON"],
                      ["物理步数", "4000", str(t["nSteps"])],
                      ["Δt / 终止 t*", "0.01 / 40", f"{t['physicalTimeStep']} / 100"],
                      ["输出间隔 / 场数", "100 / 41", f"{o['interval']} / 201"],
                      ["MPI 进程数", "16", "32（启动命令）"],
                      ["CFL / 内迭代上限", "20 / 15", f"{t['cfl']:g} / {t['maxImplicitIterations']}"],
                      ["Δτ 上限 / 松弛", "0.005 / 0.7", "0.005 / 0.7"]],
              6.90, 1.84, [2.11, 1.57, 2.09], [0.43] + [0.48] * 6, size=16)
    add_text(slide, "网格：CylinderA1.cgns，23,791 个单元，80 条圆柱壁面边；D=1，现有网格未做本次加密。\n"
             "重构：二次 Variational + CWBAP，面求积 intOrder=4；初值 (u,v,w,p)=(1,0.001,0,0)。\n"
             "物理时间采用 BDF2DualTimeLUSGS；每个 MPI 进程使用 1 个 OpenMP 线程。",
             0.65, 5.56, 12.0, 1.14, size=17.5)

    # 8. Audited numerical fields and actual residual trace.
    final = audit["final_cell_extrema"]
    def bounds(key):
        return [
            f"{value:.2e}".replace("e-0", "e-")
            if 0 < abs(value) < 0.001 else f"{value:.5g}"
            for value in final[key]
        ]
    slide = add_slide(
        prs, "已完成结果：流场数据与内迭代缺陷", str(ARCHIVE) + "\n"
        + json.dumps(audit, ensure_ascii=False, indent=2),
        "数据：归档第 4000 步 VTK-HDF 与 4000 行运行日志；统计均为无量纲单元中心值",
    )
    add_text(slide, "第 4000 步，t*=40", 0.65, 1.19, 5.8, 0.43, size=21, bold=True)
    add_text(slide, "每个物理步：内迭代前 / 后缺陷", 6.83, 1.19, 5.9, 0.43, size=21, bold=True)
    add_table(slide, [["变量", "最小值", "最大值"], ["ρ（配置常数）", "1", "1"],
                      ["u", *bounds("u")], ["v", *bounds("v")],
                      ["|V|=√(u²+v²)", *bounds("speed")], ["p（表压）", *bounds("p")]],
              0.65, 1.96, [2.30, 1.50, 1.50], [0.47] * 6, size=18)
    fit_image(slide, ASSETS / "residual_history.png", 6.18, 1.88, 6.61, 3.40)
    add_text(slide, "原生输出：Velocity（三分量）与 Pressure；ρ 不在输出场中，密度图按配置常数重建。\n"
             "已核验：41 场、0–4000 步无缺帧，全部数值有限。末步缺陷 0.3085 → 0.01065。\n"
             "各步达到 15 次内迭代上限，未达到 10⁻⁸ 阈值；本结果不等同于严格收敛或网格无关验证。",
             0.65, 5.47, 12.0, 1.31, size=18)

    # 9. Videos are truly embedded, not linked to the deleted output directory.
    slide = add_slide(
        prs, "计算过程视频：速度大小与流向速度", str(MEDIA) + "\n"
        "两段视频均嵌入PPTX。PowerPoint桌面版放映模式点击画面播放。"
        "结果来自16MPI完成的4000步计算；不是新的10000步运行。"
        "视频41个计算帧+12帧末帧停留，12fps，1280x720 H.264。",
        "视频已嵌入：0–4000 步，t*=0–40；PowerPoint 放映模式点击画面播放",
    )
    add_text(slide, "速度大小 |V|", 0.65, 1.24, 5.85, 0.44, size=22, bold=True)
    add_text(slide, "流向速度 u", 6.79, 1.24, 5.85, 0.44, size=22, bold=True)
    movie(slide, "velocity", 0.61, 1.99)
    movie(slide, "u", 6.74, 1.99)
    add_text(slide, "保留原流场色标与等值线，便于观察涡脱落、回流区和尾迹发展。\n"
             "每 100 步一个原始计算帧；12 fps 播放，末帧停留 1 秒。未生成或插值额外流场。",
             0.65, 5.81, 12.0, 0.91, size=20)

    # 10. Pressure and honest constant-density display, plus scientific scope.
    slide = add_slide(
        prs, "计算过程视频：压力与常密度", str(MEDIA) + "\n"
        "压力视频读取Pressure；密度视频由rho0=1重建，不是密度输运结果。"
        "本次保留二维层流计算，不可据此声称真实Re3900三维转捩流的定量验证。",
        "归档 Re=3900 二维常密度结果；完整配置、数据来源与方法定位见各页备注",
    )
    add_text(slide, "压力 p（表压）", 0.65, 1.24, 5.85, 0.44, size=22, bold=True)
    add_text(slide, "密度 ρ=1（模型常数）", 6.79, 1.24, 5.85, 0.44, size=22, bold=True)
    movie(slide, "pressure", 0.61, 1.99)
    movie(slide, "density", 6.74, 1.99)
    add_text(slide, "压力随涡结构演化；密度没有输运方程，因此密度画面保持均匀，不存在真实密度等值线。\n"
             "定位：用于模块验证与二维非定常可视化；定量结论还需内迭代、网格与时间步独立性检查。",
             0.65, 5.81, 12.0, 0.96, size=19)

    if len(prs.slides) != 10:
        raise ValueError("Expected exactly ten slides")
    prs.save(PPTX)
    print(PPTX)


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    audit = audit_archived_results()
    render_equations()
    build_deck(read_config(), audit)
    print(json.dumps(audit["final_cell_extrema"], ensure_ascii=False))


if __name__ == "__main__":
    main()
