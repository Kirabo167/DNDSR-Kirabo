#!/usr/bin/env python3
"""Generate a concise three-slide ACM2D cylinder case presentation.

The deck intentionally uses a plain white background and only documents the
case setup conditions needed for a short technical explanation.

Modifier: Runzhi Ma
"""

from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.collections import LineCollection
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from make_acm2d_cylinder_ppt import (
    CASE_FILE,
    ROOT,
    load_commented_json,
    load_mesh,
    segments_from_edges,
)


OUT_DIR = ROOT / "docs" / "presentations"
ASSET_DIR = OUT_DIR / "assets" / "acm2d_cylinder_case_simple"
PPTX_FILE = OUT_DIR / "ACM2D_Cylinder_Case_Setup_Simple_CN.pptx"

FONT = "Noto Sans CJK SC"
BLUE = "2F75B5"
BLUE_DARK = "1F4E79"
BLUE_LIGHT = "DDEBF7"
ORANGE = "ED7D31"
ORANGE_LIGHT = "FCE4D6"
GREEN = "70AD47"
TEXT = "222222"
GRAY = "666666"
LIGHT_GRAY = "F3F5F7"
BORDER = "D9E2F3"
WHITE = "FFFFFF"


def rgb(value: str) -> RGBColor:
    """Convert a six-digit hexadecimal color to ``RGBColor``."""

    return RGBColor.from_string(value)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = TEXT,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font: str = FONT,
):
    """Add one formatted textbox using inch-based coordinates."""

    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def add_title(slide, title: str, page: int) -> None:
    """Add a simple title, underline, and page number."""

    add_text(slide, title, 0.55, 0.34, 11.9, 0.52, size=25, color=BLUE_DARK, bold=True)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.96), Inches(12.20), Inches(0.035)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(BLUE)
    line.line.fill.background()
    add_text(slide, f"{page}/3", 12.15, 7.08, 0.55, 0.22, size=9, color=GRAY, align=PP_ALIGN.RIGHT)


def add_footer(slide) -> None:
    """Add the modifier information required by the project."""

    add_text(slide, "Runzhi Ma · DNDSR", 0.58, 7.08, 2.8, 0.22, size=9, color=GRAY)


def add_panel(slide, x, y, w, h, *, fill=WHITE, line=BORDER):
    """Add a plain rounded information panel."""

    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = rgb(fill)
    panel.line.color.rgb = rgb(line)
    panel.line.width = Pt(1.0)
    return panel


def add_lines(slide, items, x, y, w, h, *, size=15, spacing=10):
    """Add short labelled lines without decorative bullets."""

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.03)
    for idx, (label, body) in enumerate(items):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(spacing)
        label_run = paragraph.add_run()
        label_run.text = label
        label_run.font.name = FONT
        label_run.font.size = Pt(size)
        label_run.font.bold = True
        label_run.font.color.rgb = rgb(BLUE_DARK)
        body_run = paragraph.add_run()
        body_run.text = body
        body_run.font.name = FONT
        body_run.font.size = Pt(size)
        body_run.font.color.rgb = rgb(TEXT)
    return box


def add_table(slide, rows, x, y, w, h, widths, font_size=12.5):
    """Add a compact blue-header table."""

    table = slide.shapes.add_table(
        len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)
    ).table
    for col, width in enumerate(widths):
        table.columns[col].width = Inches(width)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(value)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(BLUE_LIGHT if i == 0 else WHITE)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = FONT
                    run.font.size = Pt(font_size)
                    run.font.bold = i == 0
                    run.font.color.rgb = rgb(BLUE_DARK if i == 0 else TEXT)
    return table


def plot_simple_mesh() -> Path:
    """Render the real near-cylinder mesh on a white background."""

    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    mesh = load_mesh()
    fig, ax = plt.subplots(figsize=(9.6, 4.2), dpi=180)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    colors = {"dom-1": "#4F81BD", "dom-2": "#ED7D31", "dom-3": "#A6A6A6"}
    for name, (coords, edges, boundaries) in mesh.items():
        ax.add_collection(
            LineCollection(
                segments_from_edges(coords, edges),
                colors=colors[name],
                linewidths=0.35,
                alpha=0.55,
            )
        )
        if "WALL" in boundaries:
            ax.add_collection(
                LineCollection(coords[boundaries["WALL"]], colors="#C00000", linewidths=1.6)
            )
    ax.set_xlim(-2.3, 8.0)
    ax.set_ylim(-2.1, 2.1)
    ax.set_aspect("equal", adjustable="box")
    ax.set_xlabel("x / D", fontsize=9, color="#666666")
    ax.set_ylabel("y / D", fontsize=9, color="#666666")
    ax.tick_params(labelsize=8, colors="#666666")
    ax.grid(True, color="#E7E6E6", linewidth=0.5)
    for spine in ax.spines.values():
        spine.set_color("#BFBFBF")
    ax.annotate(
        "WALL",
        xy=(0.35, 0.35),
        xytext=(1.25, 1.35),
        color="#C00000",
        fontsize=10,
        arrowprops=dict(arrowstyle="->", color="#C00000", lw=1.0),
    )
    ax.annotate(
        "尾迹加密方向",
        xy=(6.7, 0),
        xytext=(4.6, 1.4),
        color="#ED7D31",
        fontsize=10,
        arrowprops=dict(arrowstyle="->", color="#ED7D31", lw=1.0),
    )
    fig.tight_layout()
    output = ASSET_DIR / "mesh_near_white.png"
    fig.savefig(output, facecolor="white", bbox_inches="tight")
    plt.close(fig)
    return output


def build_deck(config: dict, mesh_image: Path) -> None:
    """Create the three-slide concise white-background deck."""

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "二维圆柱绕流算例设置（简版）"
    prs.core_properties.subject = "ACM二维圆柱绕流配置条件"
    prs.core_properties.author = "Runzhi Ma"
    blank = prs.slide_layouts[6]
    acm = config["acmSettings"]
    time = config["timeMarchSettings"]
    rec = config["reconstructionSettings"]
    vfv = config["vfvSettings"]

    # Slide 1: physical and geometric setup.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    add_title(slide, "二维圆柱绕流：计算域与物理条件", 1)
    add_footer(slide)
    add_panel(slide, 0.62, 1.28, 6.05, 5.25, fill=LIGHT_GRAY)
    add_text(slide, "算例基本信息", 0.95, 1.60, 2.2, 0.36, size=18, color=BLUE_DARK, bold=True)
    add_lines(
        slide,
        [
            ("求解模型：", "常密度ACM，二维稳态层流"),
            ("圆柱尺寸：", "中心 (0,0)，直径 D=1"),
            ("计算域：", "x≈[-100,100]，y≈[-99.80005,100]"),
            ("加密区域：", "x=[-2,30]，y=[-2,2]"),
            ("来流条件：", "ρ₀=1，U∞=1，p∞=0"),
            ("雷诺数：", "Re_D=20，μ=ρ₀U∞D/Re_D=0.05"),
            ("ACM参数：", "α=0.5，β²=4，人工声速尺度为2"),
        ],
        0.95,
        2.12,
        5.3,
        3.95,
        size=14.5,
        spacing=9,
    )
    # Simple flow schematic.
    add_panel(slide, 7.03, 1.28, 5.68, 5.25, fill=WHITE)
    for y in (2.40, 3.15, 3.90, 4.65):
        arrow = slide.shapes.add_shape(
            MSO_SHAPE.RIGHT_ARROW, Inches(7.62), Inches(y), Inches(4.30), Inches(0.16)
        )
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = rgb(BLUE)
        arrow.line.fill.background()
    cylinder = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(9.28), Inches(2.68), Inches(1.55), Inches(1.55)
    )
    cylinder.fill.solid()
    cylinder.fill.fore_color.rgb = rgb(WHITE)
    cylinder.line.color.rgb = rgb(ORANGE)
    cylinder.line.width = Pt(2.5)
    add_text(slide, "D=1", 9.54, 3.24, 1.02, 0.28, size=15, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "U∞=1", 7.62, 1.91, 1.25, 0.30, size=15, color=BLUE_DARK, bold=True)
    add_text(slide, "Re_D=20：稳态层流基准算例", 8.20, 5.43, 3.45, 0.40, size=16, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "用于验证压力—速度耦合、壁面条件和稳态收敛", 7.62, 5.96, 4.58, 0.35, size=12, color=GRAY, align=PP_ALIGN.CENTER)

    # Slide 2: mesh, initial field, and boundaries.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    add_title(slide, "网格、初始场与边界条件", 2)
    add_footer(slide)
    slide.shapes.add_picture(str(mesh_image), Inches(0.58), Inches(1.30), width=Inches(7.15))
    add_panel(slide, 0.62, 5.12, 7.08, 1.34, fill=LIGHT_GRAY)
    add_text(slide, "网格统计", 0.88, 5.39, 1.05, 0.30, size=14, color=BLUE_DARK, bold=True)
    add_text(
        slide,
        "23,791 单元（9,548 三角形 + 14,243 四边形）\n19,067 节点；WALL 80条边；FAR 20条边；不升阶、不二分。",
        1.86,
        5.31,
        5.46,
        0.73,
        size=12.5,
        color=TEXT,
    )
    rows = [
        ["项目", "配置与处理方式"],
        ["初始场", "[u,v,w,p]=[1,0,0,0]，全域均匀来流"],
        ["圆柱 WALL", "BCWall；静止无滑移，面速度为0"],
        ["壁面压力", "由内部外推，近似 ∂p/∂n=0"],
        ["外边界 FAR", "BCFar；入射特征取远场值，出射特征内部外推"],
        ["默认边界", "未单独命名的外边界按 BCFar 处理"],
    ]
    add_table(slide, rows, 8.00, 1.30, 4.72, 4.95, widths=[1.34, 3.38], font_size=11.5)

    # Slide 3: spatial and temporal numerical settings.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    add_title(slide, "数值方法、时间推进与并行设置", 3)
    add_footer(slide)
    spatial_rows = [
        ["空间离散", "当前设置"],
        ["重构", f"{rec['type']}，maxOrder={vfv['maxOrder']}（二次多项式）"],
        ["限制器", f"{rec['limiterType']}，enableLimiter=true"],
        ["积分阶数", f"intOrder={vfv['intOrder']}，变分迭代3次"],
        ["黎曼求解器", f"{acm['riemannSolverType']}，entropyFixRatio={acm['entropyFixRatio']:.2f}"],
        ["黏性通量", "开启，dynamicViscosity=0.05"],
    ]
    time_rows = [
        ["时间/线性求解", "当前设置"],
        ["推进方法", time["integrator"]],
        ["局部时间步", f"CFL={time['cfl']:.1f}，dt_max={time['maximumPseudoTimeStep']:.2f}"],
        ["最大步数", f"nSteps={time['nSteps']}（当前无残差自动早停）"],
        ["隐式内迭代", f"最多{time['maxImplicitIterations']}次，松弛={time['implicitRelaxation']:.1f}"],
        ["GMRES / LU-SGS", f"m={time['gmresSubspace']}，restart={time['gmresRestarts']}，LU-SGS {time['lusgsSweeps']} sweeps"],
    ]
    add_table(slide, spatial_rows, 0.62, 1.31, 6.00, 4.20, widths=[1.60, 4.40], font_size=11.6)
    add_table(slide, time_rows, 6.82, 1.31, 5.90, 4.20, widths=[1.75, 4.15], font_size=11.2)
    add_panel(slide, 0.62, 5.78, 12.10, 0.82, fill=BLUE_LIGHT, line=BLUE)
    add_text(slide, "32核运行：", 0.90, 6.02, 1.20, 0.28, size=13, color=BLUE_DARK, bold=True)
    add_text(
        slide,
        "OMP_NUM_THREADS=1 mpirun -np 32 ./app/euler.exe ../cases/acm2D/acm2D.json",
        2.05,
        6.00,
        8.45,
        0.30,
        size=12,
        color=TEXT,
        font="Noto Sans Mono CJK SC",
    )
    add_text(slide, "稳态局部伪时间", 10.52, 6.00, 1.82, 0.30, size=12, color=ORANGE, bold=True, align=PP_ALIGN.RIGHT)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(PPTX_FILE)
    print(PPTX_FILE)


def main() -> None:
    """Load the case and generate the concise deck."""

    config = load_commented_json(CASE_FILE)
    build_deck(config, plot_simple_mesh())


if __name__ == "__main__":
    main()
