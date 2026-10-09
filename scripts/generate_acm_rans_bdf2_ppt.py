#!/usr/bin/env python3
"""Build the ACM RANS/SST-BDF2 talk in the supplied 4:3 white Arial style."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs/solver-guide/acm_rans_sst_bdf2_implementation_zh.pptx"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(82, 82, 82)
WHITE = RGBColor(255, 255, 255)


def text_box(slide, text: str, x: float, y: float, w: float, h: float,
             size: float = 28, bold: bool = False, color=BLACK,
             align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.text = text
    paragraph.font.name = "Arial"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def blank_slide(presentation: Presentation, title: str, title_size: float = 44):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    text_box(slide, title, 0.5, 0.28, 9.0, 1.12, title_size)
    return slide


def bullet(slide, text: str, y: float, size: float = 29, h: float = 0.75):
    text_box(slide, "•", 0.56, y, 0.32, h, size)
    text_box(slide, text, 0.98, y, 8.45, h, size)


def outlined_box(slide, label: str, x: float, y: float, w: float, h: float,
                 size: float = 22):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = BLACK
    shape.line.width = Pt(1.3)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.04)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = label
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.name = "Arial"
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = BLACK
    return shape


def hline(slide, x: float, y: float, w: float, gray: bool = False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.012)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GRAY if gray else BLACK
    shape.line.fill.background()


def build() -> Presentation:
    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = "ACM 常密度 RANS 与 SST-BDF2 实现"
    presentation.core_properties.subject = "DNDSR ACM RANS implementation"
    presentation.core_properties.author = "DNDSR"
    presentation.core_properties.last_modified_by = "DNDSR"
    presentation.core_properties.created = datetime.now()
    presentation.core_properties.modified = datetime.now()

    # The reference has a white 4:3 canvas, a 44 pt Arial title at (0.5, 0.3),
    # and large black bullet text. Keep that visual vocabulary on every slide.
    slide = blank_slide(presentation, "ACM RANS 与 SST-BDF2 实现", 43)
    bullet(slide, "常密度四变量流场 + 独立湍流输运", 1.85, 30)
    bullet(slide, "新增 SST 的两级物理时间历史", 2.86, 30)
    bullet(slide, "二维与三维圆柱算例完成短跑验证", 3.87, 30)
    text_box(slide, "DNDSR  ·  2026-10", 0.6, 6.55, 8.7, 0.38, 19, color=GRAY)

    slide = blank_slide(presentation, "设计约束")
    bullet(slide, "流场 U = [u, v, w, p]ᵀ 保持四分量", 1.75, 28)
    bullet(slide, "湍流变量 q 独立存储与输运", 2.70, 28)
    bullet(slide, "μₜ 进入黏性通量和黏性 CFL", 3.65, 28)
    bullet(slide, "BDF2 仅开放 Laminar 与 KOmegaSST", 4.60, 28)
    text_box(slide, "保留原有 4×4 特征分解与隐式流场矩阵", 0.96, 6.25, 8.5, 0.52, 22, color=GRAY)

    slide = blank_slide(presentation, "从湍流场到流场残差")
    top = [("流场 U", 0.55, 1.80, 1.55), ("∇U", 2.65, 1.80, 1.45),
           ("闭合 μₜ", 4.62, 1.80, 1.85), ("R_flow", 7.15, 1.80, 2.0)]
    bottom = [("湍流 q", 0.55, 4.22, 1.55), ("∇q + 壁距 d", 2.65, 4.22, 1.45),
              ("R_q", 4.62, 4.22, 1.85), ("SSPRK3 / BDF2", 7.15, 4.22, 2.0)]
    for label, x, y, w in top + bottom:
        outlined_box(slide, label, x, y, w, 1.02, 21)
    for x in (2.20, 4.15, 6.60):
        text_box(slide, "→", x, 2.01, 0.40, 0.55, 27)
        text_box(slide, "→", x, 4.43, 0.40, 0.55, 27)
    text_box(slide, "↑", 5.31, 3.03, 0.5, 0.55, 31, align=PP_ALIGN.CENTER)
    text_box(slide, "U 决定对流/生产；q 决定 μₜ，再反馈流场。", 0.72, 6.36, 8.75, 0.55, 22)

    slide = blank_slide(presentation, "四种 RANS 闭合")
    text_box(slide, "模型", 0.58, 1.65, 2.55, 0.55, 23, bold=True)
    text_box(slide, "变量", 3.20, 1.65, 1.65, 0.55, 23, bold=True)
    text_box(slide, "实现要点", 5.02, 1.65, 4.25, 0.55, 23, bold=True)
    hline(slide, 0.54, 2.25, 8.80)
    rows = [
        ("SA", "ν̃", "阻尼涡黏度；壁面破坏"),
        ("Wilcox k–ω", "k, ω", "应力限制；交叉扩散"),
        ("SST k–ω", "k, ω", "F₁/F₂ 混合；近壁处理"),
        ("Realizable k–ε", "k, ε", "可实现性限制；低 Re 修正"),
    ]
    for idx, (name, variables, note) in enumerate(rows):
        y = 2.47 + idx * 1.00
        text_box(slide, name, 0.58, y, 2.55, 0.75, 23)
        text_box(slide, variables, 3.20, y, 1.65, 0.75, 23)
        text_box(slide, note, 5.02, y, 4.28, 0.75, 22)
        hline(slide, 0.54, y + 0.86, 8.80, gray=True)
    text_box(slide, "本次物理时间 BDF2 仅接入 SST；其余模型保留稳态伪时间推进。",
             0.58, 6.71, 8.85, 0.38, 18, color=GRAY)

    slide = blank_slide(presentation, "SST 的空间闭合")
    formulas = [
        "μₜ = ρ₀ a₁ k / max(a₁ω, |Ω| F₂)",
        "Dₖ = (μ + σₖ μₜ) / ρ₀",
        "Dω = (μ + σω μₜ) / ρ₀",
        "Sₖ = min(Pₖ, 20 β* kω) − β* kω",
        "Sω = γS² − βω² + 交叉扩散",
    ]
    for idx, formula in enumerate(formulas):
        text_box(slide, formula, 0.73, 1.68 + idx * 0.85, 8.70, 0.65, 27)
    hline(slide, 0.72, 6.11, 8.67, gray=True)
    text_box(slide, "F₁/F₂：连接近壁 k–ω 与远壁模型；壁面 k = 0。",
             0.73, 6.27, 8.72, 0.62, 21, color=GRAY)

    slide = blank_slide(presentation, "BE 启动与 BDF2")
    bullet(slide, "第 1 物理步：BE 系数 (1, −1, 0)", 1.65, 27)
    bullet(slide, "第 2 步起：BDF2 系数 (3/2, −2, 1/2)", 2.57, 27)
    hline(slide, 0.66, 3.54, 8.76, gray=True)
    text_box(slide, "Dₜx = (a₀xⁿ⁺¹ + a₁xⁿ + a₂xⁿ⁻¹) / Δt",
             0.76, 3.75, 8.55, 0.75, 26)
    text_box(slide, "D_U = R_U − M DₜU，M = diag(1,1,1,0)",
             0.76, 4.76, 8.55, 0.75, 26)
    text_box(slide, "D_q = R_q − Dₜq，q = (k, ω)",
             0.76, 5.76, 8.55, 0.75, 26)

    slide = blank_slide(presentation, "双时间内迭代")
    steps = [
        "固定 Uⁿ、Uⁿ⁻¹ 与 qⁿ、qⁿ⁻¹",
        "GMRES / LU-SGS 求四变量流场修正",
        "用新流场 SSPRK3 松弛 SST 时间缺陷",
        "重算 μₜ、流场缺陷和湍流缺陷",
        "两种缺陷达标才提前结束；再提交历史",
    ]
    for idx, step in enumerate(steps):
        y = 1.60 + idx * 0.99
        text_box(slide, f"{idx + 1}.", 0.61, y, 0.50, 0.72, 28)
        text_box(slide, step, 1.15, y, 8.16, 0.72, 27)
    text_box(slide, "达到内迭代上限仍可提交物理步；必须检查 converged。",
             0.65, 6.72, 8.78, 0.42, 19, color=GRAY)

    slide = blank_slide(presentation, "代码接入位置")
    files = [
        ("ACMBDF2.*", "BE/BDF2 系数与 q 时间导数"),
        ("ACMTurbulenceTransport.*", "SST 两级历史、时间缺陷、SSPRK3"),
        ("ACMSolver.hxx", "流场—SST 双时间迭代与历史提交"),
        ("ACM.cpp", "Laminar / SST 模型范围校验"),
        ("test_ACMTime.cpp", "解析导数与模型选择回归"),
    ]
    for idx, (name, role) in enumerate(files):
        y = 1.63 + idx * 1.02
        text_box(slide, name, 0.60, y, 3.88, 0.72, 20)
        text_box(slide, role, 4.55, y, 4.72, 0.72, 21)
        hline(slide, 0.60, y + 0.82, 8.75, gray=True)

    slide = blank_slide(presentation, "验证已完成")
    bullet(slide, "ACM BDF2 单测：17/17，89 个断言通过", 1.65, 27)
    bullet(slide, "二维圆柱：GMRES 与 LU-SGS 均完成两步", 2.61, 27)
    bullet(slide, "三维 Re=3900：4 rank 完成 BE → BDF2", 3.57, 27)
    bullet(slide, "最终 k、ω 有限且保持正值", 4.53, 27)
    hline(slide, 0.65, 5.62, 8.75, gray=True)
    text_box(slide, "短跑验证了运行链路，不等于时间精度或充分收敛。",
             0.71, 5.95, 8.67, 0.65, 24)

    slide = blank_slide(presentation, "当前限制与下一步")
    bullet(slide, "短跑均显示 converged = 0", 1.65, 28)
    bullet(slide, "SST 伪时间输运仍为分离式显式推进", 2.62, 28)
    bullet(slide, "BDF2 历史尚不能从 VTK-HDF 重启", 3.59, 28)
    bullet(slide, "需做内迭代、Δt 与网格收敛验证", 4.56, 28)
    text_box(slide, "物理预测还需检验壁面分辨率和升阻力统计。",
             0.97, 6.28, 8.39, 0.57, 22, color=GRAY)

    return presentation


def main() -> None:
    presentation = build()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    print(f"Wrote {OUTPUT} ({len(presentation.slides)} slides)")


if __name__ == "__main__":
    main()
