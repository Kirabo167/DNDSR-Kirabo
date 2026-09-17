#!/usr/bin/env python3
"""Generate the Chinese ACM2D cylinder-case presentation.

The presentation documents the single-file case configuration in
``cases/acm2D/acm2D.json`` and uses the real ``CylinderA1.cgns`` connectivity
to render the mesh figures.  All diagrams other than the mesh images are
native PowerPoint shapes so that the deck remains easy to edit.

Modifier: Runzhi Ma
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import h5py
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from matplotlib.collections import LineCollection
from matplotlib.patches import Circle, Rectangle
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
CASE_FILE = ROOT / "cases" / "acm2D" / "acm2D.json"
MESH_FILE = ROOT / "data" / "mesh" / "CylinderA1.cgns"
OUT_DIR = ROOT / "docs" / "presentations"
ASSET_DIR = OUT_DIR / "assets" / "acm2d_cylinder_case"
PPTX_FILE = OUT_DIR / "ACM2D_Cylinder_Case_Setup_CN.pptx"

FONT = "Noto Sans CJK SC"
FONT_MONO = "Noto Sans Mono CJK SC"

# Ensure Chinese labels embedded in the mesh PNGs use an installed CJK font.
PLOT_FONT_PATH = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"
font_manager.fontManager.addfont(PLOT_FONT_PATH)
PLOT_FONT = font_manager.FontProperties(fname=PLOT_FONT_PATH).get_name()
plt.rcParams["font.family"] = PLOT_FONT
plt.rcParams["axes.unicode_minus"] = False

# Presentation palette: dark technical background with cyan/orange accents.
NAVY = "07131F"
NAVY_2 = "0C2133"
NAVY_3 = "123047"
CYAN = "29D3C2"
CYAN_2 = "5EE8DC"
ORANGE = "FF9F43"
RED = "FF647C"
BLUE = "57A6FF"
GREEN = "6FE3A1"
WHITE = "F4F8FB"
MUTED = "A9BBC9"
GRID = "29475C"
INK = "102738"


def rgb(value: str) -> RGBColor:
    """Convert a six-digit hexadecimal color to ``RGBColor``."""

    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def load_commented_json(path: Path) -> dict:
    """Load the project's line-comment-enabled JSON configuration."""

    clean_lines = []
    for line in path.read_text(encoding="utf-8").splitlines():
        clean_lines.append(re.sub(r"\s*//.*$", "", line))
    return json.loads("\n".join(clean_lines))


def zone_edges(zone: h5py.Group) -> tuple[np.ndarray, list[tuple[int, int]], dict[str, np.ndarray]]:
    """Return XY coordinates, unique volume-cell edges, and boundary segments.

    Parameters
    ----------
    zone:
        A CGNS HDF5 zone group containing coordinates and element sections.
    """

    coords = np.column_stack(
        [
            zone["GridCoordinates"]["CoordinateX"][" data"][...],
            zone["GridCoordinates"]["CoordinateY"][" data"][...],
        ]
    )
    edges: set[tuple[int, int]] = set()
    for section, n_node in (("TriElements", 3), ("QuadElements", 4)):
        if section not in zone:
            continue
        conn = zone[section]["ElementConnectivity"][" data"][...].reshape(-1, n_node) - 1
        for cell in conn:
            for i_node in range(n_node):
                a = int(cell[i_node])
                b = int(cell[(i_node + 1) % n_node])
                edges.add((min(a, b), max(a, b)))
    boundaries = {}
    for section in ("WALL", "FAR"):
        if section in zone:
            boundaries[section] = (
                zone[section]["ElementConnectivity"][" data"][...].reshape(-1, 2) - 1
            )
    return coords, sorted(edges), boundaries


def load_mesh() -> dict[str, tuple[np.ndarray, list[tuple[int, int]], dict[str, np.ndarray]]]:
    """Load all three unstructured CGNS zones for plotting."""

    data = {}
    with h5py.File(MESH_FILE, "r") as handle:
        base = handle["Base"]
        for name in ("dom-1", "dom-2", "dom-3"):
            data[name] = zone_edges(base[name])
    return data


def segments_from_edges(coords: np.ndarray, edges: list[tuple[int, int]]) -> np.ndarray:
    """Convert edge index pairs to Matplotlib line segments."""

    idx = np.asarray(edges, dtype=int)
    return coords[idx]


def style_axis(ax: plt.Axes) -> None:
    """Apply the presentation's dark plotting style."""

    ax.set_facecolor(f"#{NAVY}")
    for spine in ax.spines.values():
        spine.set_color(f"#{GRID}")
    ax.tick_params(colors=f"#{MUTED}", labelsize=9)
    ax.grid(True, color=f"#{GRID}", alpha=0.28, linewidth=0.5)
    ax.set_xlabel("x / D", color=f"#{MUTED}", fontsize=9)
    ax.set_ylabel("y / D", color=f"#{MUTED}", fontsize=9)


def plot_mesh_assets(mesh_data: dict) -> None:
    """Render the full-domain, near-body, and wall-resolution mesh figures."""

    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    colors = {"dom-1": CYAN, "dom-2": ORANGE, "dom-3": BLUE}

    fig, ax = plt.subplots(figsize=(10.2, 5.5), dpi=180)
    fig.patch.set_facecolor(f"#{NAVY}")
    for zone_name, (coords, edges, boundaries) in mesh_data.items():
        segs = segments_from_edges(coords, edges)
        width = 0.18 if zone_name == "dom-3" else 0.32
        alpha = 0.46 if zone_name == "dom-3" else 0.68
        ax.add_collection(
            LineCollection(segs, colors=f"#{colors[zone_name]}", linewidths=width, alpha=alpha)
        )
        for boundary_name, conn in boundaries.items():
            bsegs = coords[conn]
            bcolor = RED if boundary_name == "WALL" else ORANGE
            ax.add_collection(LineCollection(bsegs, colors=f"#{bcolor}", linewidths=1.1))
    ax.add_patch(Rectangle((-2, -2), 32, 4, fill=False, ec=f"#{CYAN_2}", lw=1.3, ls="--"))
    ax.annotate(
        "近场/尾迹加密区  [-2,30] × [-2,2]",
        xy=(30, 2),
        xytext=(18, 24),
        color=f"#{CYAN_2}",
        fontsize=10,
        arrowprops=dict(arrowstyle="->", color=f"#{CYAN_2}", lw=1.0),
    )
    ax.text(-98, 88, "FAR：约 100D", color=f"#{ORANGE}", fontsize=10, weight="bold")
    ax.set_xlim(-105, 105)
    ax.set_ylim(-105, 105)
    ax.set_aspect("equal", adjustable="box")
    style_axis(ax)
    ax.set_title("CylinderA1：全计算域与三分区网格", color=f"#{WHITE}", fontsize=14, weight="bold")
    fig.tight_layout()
    fig.savefig(ASSET_DIR / "mesh_full.png", facecolor=fig.get_facecolor(), bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(10.2, 4.6), dpi=200)
    fig.patch.set_facecolor(f"#{NAVY}")
    for zone_name, (coords, edges, boundaries) in mesh_data.items():
        segs = segments_from_edges(coords, edges)
        ax.add_collection(
            LineCollection(segs, colors=f"#{colors[zone_name]}", linewidths=0.35, alpha=0.65)
        )
        if "WALL" in boundaries:
            ax.add_collection(
                LineCollection(coords[boundaries["WALL"]], colors=f"#{RED}", linewidths=1.7)
            )
    ax.add_patch(Circle((0, 0), 0.5, facecolor=f"#{NAVY_2}", edgecolor=f"#{RED}", lw=1.5))
    ax.annotate(
        "WALL：80 条线性边",
        xy=(0.35, 0.35),
        xytext=(1.3, 1.45),
        color=f"#{RED}",
        fontsize=10,
        arrowprops=dict(arrowstyle="->", color=f"#{RED}", lw=1.0),
    )
    ax.annotate(
        "尾迹方向",
        xy=(7.2, 0),
        xytext=(4.8, 1.55),
        color=f"#{ORANGE}",
        fontsize=10,
        arrowprops=dict(arrowstyle="->", color=f"#{ORANGE}", lw=1.0),
    )
    ax.set_xlim(-2.4, 8.2)
    ax.set_ylim(-2.15, 2.15)
    ax.set_aspect("equal", adjustable="box")
    style_axis(ax)
    ax.set_title("圆柱近场与尾迹网格", color=f"#{WHITE}", fontsize=14, weight="bold")
    fig.tight_layout()
    fig.savefig(ASSET_DIR / "mesh_near.png", facecolor=fig.get_facecolor(), bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(5.4, 5.4), dpi=200)
    fig.patch.set_facecolor(f"#{NAVY}")
    coords, edges, boundaries = mesh_data["dom-1"]
    ax.add_collection(
        LineCollection(segments_from_edges(coords, edges), colors=f"#{CYAN}", linewidths=0.55, alpha=0.72)
    )
    ax.add_collection(LineCollection(coords[boundaries["WALL"]], colors=f"#{RED}", linewidths=2.0))
    ax.add_patch(Circle((0, 0), 0.5, facecolor=f"#{NAVY_2}", edgecolor=f"#{RED}", lw=1.4))
    ax.set_xlim(-0.9, 0.9)
    ax.set_ylim(-0.9, 0.9)
    ax.set_aspect("equal", adjustable="box")
    style_axis(ax)
    ax.set_title("圆柱壁面局部网格", color=f"#{WHITE}", fontsize=14, weight="bold")
    fig.tight_layout()
    fig.savefig(ASSET_DIR / "mesh_wall.png", facecolor=fig.get_facecolor(), bbox_inches="tight")
    plt.close(fig)


def set_shape_fill(shape, color: str) -> None:
    """Set a solid fill on a PowerPoint shape."""

    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)


def set_shape_line(shape, color: str, width: float = 1.0) -> None:
    """Set a shape outline color and width in points."""

    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = WHITE,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    font: str = FONT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.03,
):
    """Add a formatted textbox using dimensions in inches."""

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rich_lines(slide, lines: list[tuple[str, str]], x, y, w, h, size=15.5):
    """Add a list of title/body lines with consistent spacing."""

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.04)
    for i, (title, body) in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_after = Pt(10)
        p.level = 0
        r = p.add_run()
        r.text = title
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = rgb(CYAN_2)
        r = p.add_run()
        r.text = body
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = rgb(WHITE)
    return box


def add_card(slide, x, y, w, h, *, fill=NAVY_2, line=GRID, radius=True):
    """Add a dark information card and return the shape."""

    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    card = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    set_shape_fill(card, fill)
    set_shape_line(card, line, 0.9)
    return card


def add_badge(slide, text, x, y, w, *, fill=CYAN, color=NAVY, size=11):
    """Add a compact rounded badge."""

    badge = add_card(slide, x, y, w, 0.34, fill=fill, line=fill)
    add_text(
        slide,
        text,
        x,
        y + 0.005,
        w,
        0.30,
        size=size,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    return badge


def add_header(slide, number: int, title: str, kicker: str = "ACM2D · CYLINDERA1") -> None:
    """Add the recurring slide header and accent rule."""

    add_text(slide, f"{number:02d}", 0.45, 0.30, 0.55, 0.38, size=14, color=CYAN, bold=True)
    add_text(slide, kicker, 1.02, 0.31, 3.6, 0.34, size=10.5, color=MUTED, bold=True)
    add_text(slide, title, 0.48, 0.72, 12.0, 0.58, size=25, color=WHITE, bold=True)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.48), Inches(1.36), Inches(12.35), Inches(0.035))
    set_shape_fill(rule, GRID)
    rule.line.fill.background()
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.48), Inches(1.36), Inches(1.55), Inches(0.055))
    set_shape_fill(accent, CYAN)
    accent.line.fill.background()


def add_footer(slide, page: int) -> None:
    """Add author, project, and page number."""

    add_text(slide, "Runzhi Ma  ·  DNDSR", 0.50, 7.12, 3.5, 0.22, size=8.5, color=MUTED)
    add_text(slide, f"{page:02d}", 12.20, 7.10, 0.58, 0.24, size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def set_slide_background(slide, color=NAVY) -> None:
    """Set a full-slide solid background."""

    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(color)


def add_stat(slide, value, label, x, y, w, *, accent=CYAN):
    """Add a compact numeric statistic card."""

    add_card(slide, x, y, w, 1.05, fill=NAVY_2, line=GRID)
    add_text(slide, value, x + 0.12, y + 0.12, w - 0.24, 0.43, size=21, color=accent, bold=True)
    add_text(slide, label, x + 0.12, y + 0.60, w - 0.24, 0.28, size=10.5, color=MUTED)


def add_table(slide, rows, x, y, w, h, widths=None, font_size=12.5):
    """Add a styled table from a rectangular list of strings."""

    nrows = len(rows)
    ncols = len(rows[0])
    table = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(h)).table
    if widths:
        for idx, width in enumerate(widths):
            table.columns[idx].width = Inches(width)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(value)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.03)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(NAVY_3 if i == 0 else NAVY_2)
            cell.border_left = None if hasattr(cell, "border_left") else None
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.name = FONT
                    run.font.size = Pt(font_size if i else font_size - 0.5)
                    run.font.bold = i == 0
                    run.font.color.rgb = rgb(CYAN_2 if i == 0 else WHITE)
    return table


def add_arrow_between(slide, x1, y1, x2, y2, color=CYAN, width=1.6):
    """Add a simple connector line between two slide coordinates."""

    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def build_presentation(config: dict) -> None:
    """Build and save the complete editable PowerPoint deck."""

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "二维圆柱绕流算例设置"
    prs.core_properties.subject = "常密度ACM二维圆柱绕流配置说明"
    prs.core_properties.author = "Runzhi Ma"
    prs.core_properties.keywords = "ACM, CFD, Cylinder, DNDSR, GMRES, CWBAP"
    blank = prs.slide_layouts[6]

    acm = config["acmSettings"]
    time = config["timeMarchSettings"]
    rec = config["reconstructionSettings"]
    vfv = config["vfvSettings"]

    # Slide 1: cover.
    slide = prs.slides.add_slide(blank)
    set_slide_background(slide)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5))
    set_shape_fill(accent, CYAN)
    accent.line.fill.background()
    add_badge(slide, "CASE SETUP", 0.72, 0.72, 1.48)
    add_text(slide, "二维圆柱绕流\n算例设置", 0.72, 1.36, 6.2, 1.65, size=38, bold=True)
    add_text(
        slide,
        "常密度人工可压缩性方法（ACM）",
        0.75,
        3.17,
        5.8,
        0.44,
        size=18,
        color=CYAN_2,
        bold=True,
    )
    add_text(
        slide,
        "Re_D = 20  ·  二次变分重构  ·  CWBAP  ·  Roe  ·  GMRES / LU-SGS",
        0.75,
        3.75,
        6.8,
        0.45,
        size=13.5,
        color=MUTED,
    )
    add_text(slide, "CylinderA1.cgns  |  23,791 cells  |  32 MPI ranks", 0.75, 4.30, 6.5, 0.34, size=12, color=WHITE)
    # Native vector cylinder-flow hero.
    for i, offset in enumerate((-0.82, -0.42, 0.0, 0.42, 0.82)):
        y = 3.45 + offset
        arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(7.40), Inches(y), Inches(4.65), Inches(0.13))
        set_shape_fill(arrow, CYAN if i in (1, 2, 3) else BLUE)
        arrow.line.fill.background()
    cylinder = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.06), Inches(2.56), Inches(1.78), Inches(1.78))
    set_shape_fill(cylinder, NAVY_3)
    set_shape_line(cylinder, RED, 3.0)
    add_text(slide, "D = 1", 9.25, 3.16, 1.4, 0.30, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "U∞ = 1", 7.43, 2.12, 1.4, 0.32, size=13, color=CYAN_2, bold=True)
    add_text(slide, "Runzhi Ma", 0.75, 6.55, 2.6, 0.36, size=12, color=WHITE, bold=True)
    add_text(slide, "DNDSR · 2026", 0.75, 6.91, 2.6, 0.28, size=10, color=MUTED)

    # Slide 2: case positioning.
    slide = prs.slides.add_slide(blank)
    set_slide_background(slide)
    add_header(slide, 2, "算例定位：先建立可验证的稳态层流基线")
    add_footer(slide, 2)
    cards = [
        ("01", "物理目标", "验证常密度ACM在不可压黏性绕流中的压力—速度耦合与壁面处理。", CYAN),
        ("02", "工况选择", "采用 Re_D=20 的稳态层流，避免把物理非定常涡脱落误当作数值不收敛。", ORANGE),
        ("03", "数值目标", "检查二次VR、CWBAP、一般α Roe通量以及隐式GMRES/LU-SGS组合。", BLUE),
    ]
    for i, (num, title, body, color) in enumerate(cards):
        x = 0.55 + i * 4.22
        add_card(slide, x, 1.75, 3.85, 2.22, fill=NAVY_2, line=color)
        add_text(slide, num, x + 0.22, 1.98, 0.55, 0.42, size=17, color=color, bold=True)
        add_text(slide, title, x + 0.82, 1.96, 2.65, 0.44, size=18, bold=True)
        add_text(slide, body, x + 0.25, 2.60, 3.30, 1.05, size=13.5, color=MUTED)
    add_card(slide, 0.55, 4.35, 12.05, 1.83, fill=NAVY_3, line=GRID)
    add_badge(slide, "核心假设", 0.84, 4.68, 1.28, fill=ORANGE)
    add_rich_lines(
        slide,
        [
            ("常密度：", "ρ₀=1，守恒/状态变量按 [u, v, w, p] 组织；二维中 w=0。"),
            ("稳态策略：", "通过局部伪时间推进寻找定常解；该时间变量不代表真实物理时间。"),
            ("适用范围：", "用于算法基线和稳态阻力/压力分布验证，不直接用于 Re=100 涡脱落频率计算。"),
        ],
        2.35,
        4.60,
        9.85,
        1.35,
        size=13.2,
    )

    # Slide 3: global mesh.
    slide = prs.slides.add_slide(blank)
    set_slide_background(slide)
    add_header(slide, 3, "计算域与网格：远场约 100D，尾迹区域定向加密")
    add_footer(slide, 3)
    slide.shapes.add_picture(str(ASSET_DIR / "mesh_full.png"), Inches(0.55), Inches(1.55), width=Inches(5.18))
    add_stat(slide, "23,791", "总体单元数", 6.08, 1.73, 2.82, accent=CYAN)
    add_stat(slide, "19,067", "总体节点数", 9.20, 1.73, 2.82, accent=BLUE)
    add_stat(slide, "9,548", "三角形", 6.08, 3.04, 2.82, accent=ORANGE)
    add_stat(slide, "14,243", "四边形", 9.20, 3.04, 2.82, accent=GREEN)
    add_card(slide, 6.08, 4.40, 5.94, 1.72, fill=NAVY_2, line=GRID)
    add_text(slide, "计算域范围", 6.38, 4.66, 1.22, 0.32, size=12, color=CYAN_2, bold=True)
    add_text(slide, "x ≈ [-100, 100]   ·   y ≈ [-99.80005, 100]", 6.38, 5.05, 5.15, 0.34, size=15, color=WHITE, bold=True)
    add_text(slide, "近场/尾迹块：[-2,30] × [-2,2]；外边界距圆柱约 100D", 6.38, 5.57, 5.15, 0.28, size=11, color=MUTED)

    # Slide 4: near mesh and geometry.
    slide = prs.slides.add_slide(blank)
    set_slide_background(slide)
    add_header(slide, 4, "圆柱近场：D=1，壁面由 80 条线性边离散")
    add_footer(slide, 4)
    slide.shapes.add_picture(str(ASSET_DIR / "mesh_near.png"), Inches(0.55), Inches(1.66), width=Inches(7.62))
    slide.shapes.add_picture(str(ASSET_DIR / "mesh_wall.png"), Inches(8.38), Inches(1.66), width=Inches(3.93))
    add_card(slide, 0.68, 5.72, 11.70, 0.92, fill=NAVY_3, line=ORANGE)
    add_text(slide, "几何精度提示", 0.92, 5.96, 1.45, 0.29, size=12, color=ORANGE, bold=True)
    add_text(
        slide,
        "meshElevation=0、meshDirectBisect=0：流场采用高阶重构，但圆柱边界仍是80段直线近似；高精度阻力评估需单独检查曲率/网格收敛。",
        2.30,
        5.88,
        9.68,
        0.53,
        size=12.7,
        color=WHITE,
    )

    # Slide 5: physical parameters.
    slide = prs.slides.add_slide(blank)
    set_slide_background(slide)
    add_header(slide, 5, "无量纲参数与ACM模型参数")
    add_footer(slide, 5)
    add_card(slide, 0.55, 1.68, 4.12, 1.62, fill=NAVY_3, line=CYAN)
    add_text(slide, "Re_D = ρ₀ U∞ D / μ = 20", 0.86, 2.02, 3.55, 0.48, size=22, color=CYAN_2, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "ρ₀ = U∞ = D = 1  ⇒  μ = 0.05", 0.86, 2.58, 3.55, 0.34, size=14, color=WHITE, align=PP_ALIGN.CENTER)
    add_card(slide, 0.55, 3.58, 4.12, 1.62, fill=NAVY_3, line=ORANGE)
    add_text(slide, "cₐ = √(β² / ρ₀) = 2", 0.86, 3.93, 3.55, 0.48, size=22, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "β²=4 控制人工压力波传播尺度", 0.86, 4.49, 3.55, 0.34, size=13, color=WHITE, align=PP_ALIGN.CENTER)
    rows = [
        ["JSON参数", "取值", "选择说明"],
        ["rho0", f"{acm['rho0']:.1f}", "常密度基准"],
        ["alpha", f"{acm['alpha']:.1f}", "启用一般 α 特征系统"],
        ["beta2", f"{acm['beta2']:.1f}", "人工声速尺度 cₐ=2"],
        ["dynamicViscosity", f"{acm['dynamicViscosity']:.2f}", "对应 Re_D=20"],
        ["riemannSolverType", acm["riemannSolverType"], "一般 α Roe耗散"],
        ["entropyFixRatio", f"{acm['entropyFixRatio']:.2f}", "特征值近零熵修正"],
        ["enableViscousFlux", "true", "计算不可压黏性通量"],
    ]
    add_table(slide, rows, 4.98, 1.68, 7.72, 4.86, widths=[2.15, 1.15, 4.42], font_size=11.6)

    # Slide 6: initial and boundary conditions.
    slide = prs.slides.add_slide(blank)
    set_slide_background(slide)
    add_header(slide, 6, "初始场与边界条件：壁面无滑移，远场按特征波处理")
    add_footer(slide, 6)
    add_card(slide, 0.55, 1.66, 7.25, 4.80, fill=NAVY_2, line=GRID)
    for y in (2.48, 3.18, 3.88, 4.58):
        arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(1.03), Inches(y), Inches(5.75), Inches(0.18))
        set_shape_fill(arrow, CYAN if y in (3.18, 3.88) else BLUE)
        arrow.line.fill.background()
    cylinder = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3.63), Inches(2.69), Inches(1.45), Inches(1.45))
    set_shape_fill(cylinder, NAVY_3)
    set_shape_line(cylinder, RED, 2.6)
    add_text(slide, "WALL", 3.75, 3.23, 1.2, 0.28, size=13, bold=True, align=PP_ALIGN.CENTER)
    add_badge(slide, "FAR  [1,0,0,0]", 0.94, 1.98, 1.85, fill=BLUE, size=10)
    add_text(slide, "来流 U∞=1", 1.02, 5.30, 1.6, 0.30, size=12, color=CYAN_2, bold=True)
    add_text(slide, "出射特征：内部外推", 5.42, 2.05, 1.95, 0.28, size=11.5, color=ORANGE, bold=True)
    add_text(slide, "入射特征：远场给定", 5.42, 5.38, 1.95, 0.28, size=11.5, color=BLUE, bold=True)
    add_card(slide, 8.05, 1.66, 4.65, 1.21, fill=NAVY_3, line=CYAN)
    add_text(slide, "初始场", 8.30, 1.90, 0.92, 0.30, size=13, color=CYAN_2, bold=True)
    add_text(slide, "[u,v,w,p] = [1,0,0,0]", 9.16, 1.88, 3.02, 0.34, size=15, bold=True)
    add_text(slide, "全域均匀来流，减少远场启动扰动", 8.30, 2.34, 3.92, 0.24, size=10.5, color=MUTED)
    add_card(slide, 8.05, 3.11, 4.65, 1.37, fill=NAVY_3, line=RED)
    add_text(slide, "BCWall", 8.30, 3.37, 1.0, 0.30, size=13, color=RED, bold=True)
    add_text(slide, "u_g = 2u_wall − u_i", 9.28, 3.35, 2.82, 0.32, size=14, bold=True)
    add_text(slide, "面速度为0；压力内部外推，近似 ∂p/∂n=0", 8.30, 3.87, 3.96, 0.32, size=10.8, color=MUTED)
    add_card(slide, 8.05, 4.72, 4.65, 1.58, fill=NAVY_3, line=BLUE)
    add_text(slide, "BCFar", 8.30, 4.98, 0.9, 0.30, size=13, color=BLUE, bold=True)
    add_text(slide, "一般 α 特征边界", 9.15, 4.96, 2.92, 0.32, size=14, bold=True)
    add_text(slide, "入射波采用远场状态；出射波由内部解外推。\n未命名外边界也回退为 BCFar。", 8.30, 5.43, 3.95, 0.58, size=10.8, color=MUTED)

    # Slide 7: reconstruction pipeline.
    slide = prs.slides.add_slide(blank)
    set_slide_background(slide)
    add_header(slide, 7, "空间离散：二次变分重构 + ACM特征CWBAP + Roe通量")
    add_footer(slide, 7)
    stages = [
        ("单元平均量", "[u,v,w,p]", BLUE),
        ("变分重构", "p=2 · 迭代3次", CYAN),
        ("CWBAP限制", "一般α特征空间", ORANGE),
        ("面上积分", "intOrder=4", GREEN),
        ("数值通量", "Roe + 黏性", RED),
    ]
    for i, (title, body, color) in enumerate(stages):
        x = 0.52 + i * 2.53
        add_card(slide, x, 2.02, 2.05, 1.45, fill=NAVY_2, line=color)
        add_text(slide, f"{i+1}", x + 0.15, 2.20, 0.36, 0.30, size=12, color=color, bold=True)
        add_text(slide, title, x + 0.52, 2.17, 1.32, 0.34, size=14, bold=True)
        add_text(slide, body, x + 0.16, 2.73, 1.73, 0.40, size=11.5, color=MUTED, align=PP_ALIGN.CENTER)
        if i < len(stages) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x + 2.09), Inches(2.56), Inches(0.36), Inches(0.34))
            set_shape_fill(arrow, GRID)
            arrow.line.fill.background()
    add_card(slide, 0.55, 4.00, 12.03, 1.79, fill=NAVY_3, line=GRID)
    add_rich_lines(
        slide,
        [
            ("重构阶数：", f"maxOrder={vfv['maxOrder']} 表示保留至二次多项式；实际空间收敛阶需要网格加密试验确认。"),
            ("限制器：", f"{rec['limiterType']} 在ACM一般α左右特征矩阵下限制多项式系数；WBAP_nStd={vfv['WBAP_nStd']:.0f}，normWBAP=false。"),
            ("积分与边界：", f"intOrder={vfv['intOrder']}，bcWeight={vfv['bcWeight']:.1f}；壁面/远场状态参与边界面重构和通量计算。"),
        ],
        0.86,
        4.27,
        11.25,
        1.25,
        size=12.5,
    )

    # Slide 8: time marching.
    slide = prs.slides.add_slide(blank)
    set_slide_background(slide)
    add_header(slide, 8, "时间推进：局部伪时间下的隐式Euler—GMRES—LU-SGS")
    add_footer(slide, 8)
    labels = [
        ("高阶残差 R(Q)", CYAN),
        ("冻结一阶隐式算子", BLUE),
        ("GMRES", ORANGE),
        ("4×4 LU-SGS预条件", GREEN),
        ("松弛更新 Q", RED),
    ]
    for i, (label, color) in enumerate(labels):
        x = 0.60 + i * 2.49
        add_card(slide, x, 1.85, 2.02, 1.07, fill=NAVY_2, line=color)
        add_text(slide, label, x + 0.14, 2.14, 1.74, 0.46, size=12.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if i < len(labels) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.04), Inches(2.20), Inches(0.38), Inches(0.24))
            set_shape_fill(arrow, GRID)
            arrow.line.fill.background()
    rows = [
        ["设置", "取值", "作用"],
        ["integrator", time["integrator"], "后向Euler隐式伪时间"],
        ["nSteps", str(time["nSteps"]), "固定外层步数；当前无自动早停"],
        ["CFL / dt_max", f"{time['cfl']:.1f} / {time['maximumPseudoTimeStep']:.2f}", "局部谱半径步长并设置上限"],
        ["maxImplicitIterations", str(time["maxImplicitIterations"]), "每个外层步最多5次非线性修正"],
        ["relaxation", f"{time['implicitRelaxation']:.1f}", "抑制过大隐式更新"],
        ["GMRES", f"m={time['gmresSubspace']}, restart={time['gmresRestarts']}", f"相对容差 {time['gmresRelativeTolerance']:.0e}"],
        ["LU-SGS", f"{time['lusgsSweeps']} sweeps", "每次预条件应用扫描2次"],
    ]
    add_table(slide, rows, 0.60, 3.37, 8.05, 3.15, widths=[2.05, 2.28, 3.72], font_size=10.5)
    add_card(slide, 8.94, 3.37, 3.68, 3.15, fill=NAVY_3, line=ORANGE)
    add_badge(slide, "重要解释", 9.25, 3.71, 1.15, fill=ORANGE, size=10.5)
    add_text(slide, "这是稳态求解器的\n伪时间推进", 9.24, 4.28, 3.02, 0.78, size=19, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "局部时间步可加速收敛，\n但不能用于提取真实频率或Strouhal数。", 9.35, 5.22, 2.82, 0.70, size=12, color=MUTED, align=PP_ALIGN.CENTER)

    # Slide 9: parallel setup.
    slide = prs.slides.add_slide(blank)
    set_slide_background(slide)
    add_header(slide, 9, "并行计算：METIS KWAY划分，32个MPI进程运行")
    add_footer(slide, 9)
    add_card(slide, 0.55, 1.68, 6.15, 3.78, fill=NAVY_2, line=GRID)
    palette = [CYAN, BLUE, ORANGE, GREEN, RED, CYAN_2, "A478FF", "E4C15A"]
    for rank in range(32):
        col = rank % 8
        row = rank // 8
        x = 0.86 + col * 0.68
        y = 2.05 + row * 0.68
        tile = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(x), Inches(y), Inches(0.54), Inches(0.48))
        set_shape_fill(tile, palette[rank % len(palette)])
        tile.line.fill.background()
        add_text(slide, f"{rank}", x + 0.04, y + 0.095, 0.46, 0.18, size=7.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "32 MPI ranks（示意）", 1.06, 4.91, 2.4, 0.30, size=12, color=CYAN_2, bold=True)
    add_text(slide, "实际分区由单元邻接图决定，并非规则块状切分", 3.06, 4.91, 3.1, 0.30, size=10.5, color=MUTED)
    add_card(slide, 6.98, 1.68, 5.63, 1.68, fill=NAVY_3, line=CYAN)
    add_text(slide, "启动命令", 7.27, 1.96, 1.1, 0.28, size=12, color=CYAN_2, bold=True)
    command = "OMP_NUM_THREADS=1 mpirun -np 32 \\\n  ./app/euler.exe ../cases/acm2D/acm2D.json"
    add_text(slide, command, 7.27, 2.37, 4.92, 0.66, size=10.6, color=WHITE, font=FONT_MONO)
    add_card(slide, 6.98, 3.65, 5.63, 1.81, fill=NAVY_3, line=GRID)
    add_rich_lines(
        slide,
        [
            ("划分：", "metisType=KWAY，metisUfactor=20，metisNcuts=3。"),
            ("线程：", "每个MPI进程设置 OMP_NUM_THREADS=1，避免32×多线程过度占用。"),
            ("配置位置：", "MPI进程数由启动命令指定，不写入JSON。"),
        ],
        7.22,
        3.93,
        5.03,
        1.28,
        size=11.3,
    )

    # Slide 10: validation and limitations.
    slide = prs.slides.add_slide(blank)
    set_slide_background(slide)
    add_header(slide, 10, "冒烟验证结果与当前使用边界")
    add_footer(slide, 10)
    add_card(slide, 0.55, 1.70, 5.88, 2.36, fill=NAVY_3, line=GREEN)
    add_badge(slide, "2 MPI · 1 implicit step", 0.84, 2.00, 2.05, fill=GREEN, size=10.5)
    add_text(slide, "8.695 × 10²", 0.90, 2.64, 2.10, 0.52, size=24, color=RED, bold=True, align=PP_ALIGN.CENTER)
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(3.03), Inches(2.76), Inches(0.72), Inches(0.30))
    set_shape_fill(arrow, CYAN)
    arrow.line.fill.background()
    add_text(slide, "2.959 × 10⁻¹", 3.77, 2.64, 2.22, 0.52, size=24, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "初始非线性残差", 1.11, 3.24, 1.74, 0.28, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "单步后残差", 4.09, 3.24, 1.58, 0.28, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_card(slide, 0.55, 4.36, 5.88, 1.52, fill=NAVY_2, line=GRID)
    add_rich_lines(
        slide,
        [
            ("已验证：", "网格读取与分区、边界匹配、二次VR/CWBAP、Roe、GMRES/LU-SGS均正常进入计算。"),
            ("未验证：", "尚未完成5000步收敛、阻力系数与压力分布对参考数据的定量对比。"),
        ],
        0.82,
        4.66,
        5.34,
        0.99,
        size=11.4,
    )
    add_card(slide, 6.75, 1.70, 5.85, 4.18, fill=NAVY_2, line=ORANGE)
    add_text(slide, "使用前必须知道", 7.08, 2.02, 2.25, 0.34, size=16, color=ORANGE, bold=True)
    cautions = [
        "nSteps=5000 为固定执行步数，当前驱动没有按残差自动提前终止。",
        "当前ACM初版尚未接通 solution / restart / VTK 输出。",
        "隐式算子使用冻结的一阶面线性化，高阶重构保留在非线性残差中。",
        "若目标改为 Re=100 涡脱落，应使用真实物理时间推进并加入非对称扰动与时序输出。",
    ]
    for i, item in enumerate(cautions):
        y = 2.62 + i * 0.70
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.08), Inches(y + 0.07), Inches(0.13), Inches(0.13))
        set_shape_fill(dot, ORANGE if i < 3 else RED)
        dot.line.fill.background()
        add_text(slide, item, 7.34, y, 4.76, 0.54, size=11.7, color=WHITE)

    # Slide 11: JSON navigation and handoff.
    slide = prs.slides.add_slide(blank)
    set_slide_background(slide)
    add_header(slide, 11, "单文件配置导航与后续调参入口")
    add_footer(slide, 11)
    rows = [
        ["JSON模块", "当前职责", "优先调节内容"],
        ["acmSettings", "物理模型、α/β²、黏性和Roe通量", "Re、α、β²、熵修正"],
        ["timeMarchSettings", "伪时间步与隐式线性求解", "CFL、dt_max、nSteps、GMRES"],
        ["turbulenceSettings", "统一模型接口；本算例为Laminar", "Re升高后再选择SA/SST等"],
        ["meshSettings", "CGNS读取和MPI划分", "网格文件、METIS参数"],
        ["reconstructionSettings", "重构类型和限制器", "CWBAP开关/类型"],
        ["vfvSettings", "多项式次数、积分与VR权重", "maxOrder、intOrder、WBAP参数"],
        ["boundaryConditions", "按CGNS名称映射WALL/FAR", "壁速、远场速度/压力"],
        ["initialState", "全域初始化", "[u,v,w,p]"],
    ]
    add_table(slide, rows, 0.55, 1.69, 8.58, 4.86, widths=[2.05, 3.15, 3.38], font_size=10.7)
    add_card(slide, 9.45, 1.69, 3.15, 2.12, fill=NAVY_3, line=CYAN)
    add_text(slide, "唯一算例文件", 9.74, 1.99, 2.58, 0.34, size=15, color=CYAN_2, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "cases/acm2D/\nacm2D.json", 9.74, 2.53, 2.58, 0.72, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER, font=FONT_MONO)
    add_card(slide, 9.45, 4.13, 3.15, 2.42, fill=NAVY_3, line=ORANGE)
    add_text(slide, "推荐调整顺序", 9.74, 4.43, 2.58, 0.34, size=15, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "① 先跑一阶/小CFL\n② 再启用二次+CWBAP\n③ 最后提高CFL并检查阻力", 9.76, 5.00, 2.54, 1.08, size=12.3, color=WHITE, align=PP_ALIGN.CENTER)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(PPTX_FILE)


def main() -> None:
    """Generate mesh assets and the PowerPoint presentation."""

    config = load_commented_json(CASE_FILE)
    mesh_data = load_mesh()
    plot_mesh_assets(mesh_data)
    build_presentation(config)
    print(PPTX_FILE)


if __name__ == "__main__":
    main()
